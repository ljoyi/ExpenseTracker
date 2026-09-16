from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "记账App_MVP功能清单.md"
OUTPUT = ROOT / "记账App_MVP功能表.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
P0_FILL = PatternFill("solid", fgColor="FCE4D6")
P1_FILL = PatternFill("solid", fgColor="FFF2CC")
P2_FILL = PatternFill("solid", fgColor="E2F0D9")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
THIN_BORDER = Border(
    left=Side(style="thin", color="B7C9D6"),
    right=Side(style="thin", color="B7C9D6"),
    top=Side(style="thin", color="B7C9D6"),
    bottom=Side(style="thin", color="B7C9D6"),
)


def read_table(heading: str) -> list[list[str]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = lines.index(heading) + 1
    table_lines: list[str] = []

    for line in lines[start:]:
        if table_lines and not line.startswith("|"):
            break
        if line.startswith("|"):
            table_lines.append(line)

    rows: list[list[str]] = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def style_table(ws, widths: list[int], freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def build_feature_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "MVP功能总表"
    rows = read_table("## 4. MVP 功能表单")
    for row in rows:
        ws.append(row)

    style_table(ws, [8, 14, 18, 9, 44, 58, 10])
    priority_column = rows[0].index("优先级") + 1
    for row_number in range(2, ws.max_row + 1):
        priority = ws.cell(row_number, priority_column).value
        fill = {"P0": P0_FILL, "P1": P1_FILL, "P2": P2_FILL}.get(priority)
        if fill:
            ws.cell(row_number, priority_column).fill = fill
            ws.cell(row_number, priority_column).font = Font(bold=True)
        ws.row_dimensions[row_number].height = 38


def build_page_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("MVP页面清单")
    rows = read_table("## 5. MVP 页面清单")
    for row in rows:
        ws.append(row)
    style_table(ws, [18, 52, 12])
    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 34


def append_titled_list(ws, title: str, items: list[str], row: int) -> int:
    ws.cell(row, 1, title).font = Font(bold=True, size=12, color="1F4E78")
    ws.cell(row, 1).fill = SECTION_FILL
    ws.cell(row, 1).border = THIN_BORDER
    ws.cell(row, 2).fill = SECTION_FILL
    ws.cell(row, 2).border = THIN_BORDER
    row += 1
    for item in items:
        ws.cell(row, 1, "-")
        ws.cell(row, 2, item)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, 1).border = THIN_BORDER
        ws.cell(row, 2).border = THIN_BORDER
        row += 1
    return row + 1


def build_scope_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("范围与验收")
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 92

    row = 1
    row = append_titled_list(
        ws,
        "首发目标",
        [
            "默认进入账单页，用户无需登录即可开始记账。",
            "用户可在 10 秒内完成一笔普通支出记录。",
            "用户可按日期查看、修改和删除账单。",
            "用户可查看本月及任意月份的收支概要、分类占比和每日趋势。",
            "收入与支出使用统一的数据模型，统计准确。",
        ],
        row,
    )
    row = append_titled_list(
        ws,
        "本期不做",
        [
            "多用户、多人协作、账本成员",
            "多账本切换",
            "资产总览、账户余额联动",
            "预算与超支提醒",
            "银行账单导入、自动记账",
            "云同步、跨设备同步、登录体系",
            "定期账单、报销流程、复杂退款流程",
            "多币种和汇率换算",
            "Widget、Apple Watch、快捷指令",
        ],
        row,
    )
    row = append_titled_list(
        ws,
        "MVP验收指标",
        [
            "新用户首次启动后，可在 3 次点击内到达记账输入状态。",
            "普通支出从打开记账页到保存完成，目标不超过 10 秒。",
            "新增、修改、删除账单后，列表、汇总和统计结果同步且一致。",
            "切换月份、日期和统计周期时，不出现串月、漏账或重复统计。",
            "App 冷启动无崩溃，离线可完成全部核心功能。",
            "使用 1,000 条账单进行基本滑动、筛选和统计时保持可用。",
        ],
        row,
    )
    row = append_titled_list(
        ws,
        "推荐开发顺序",
        [
            "数据模型与本地存储：账单、分类、账本。",
            "记账闭环：新增、保存、账单列表、详情。",
            "管理闭环：修改、删除、月份切换、日历筛选。",
            "洞察闭环：统计汇总、分类占比、趋势图、报表。",
            "体验完善：分类管理、空状态、基础设置、连续记账。",
            "稳定性验证：汇总一致性、边界金额、日期时区、性能测试。",
        ],
        row,
    )


def main() -> None:
    wb = Workbook()
    build_feature_sheet(wb)
    build_page_sheet(wb)
    build_scope_sheet(wb)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
